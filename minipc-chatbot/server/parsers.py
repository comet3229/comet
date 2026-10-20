"""문서(PDF/PPTX/DOCX/TXT/MD) → 텍스트 청크 변환. 파일은 메모리에서만 처리하고 저장하지 않는다."""
import io
from pathlib import PurePath

CHUNK_SIZE = 800
CHUNK_OVERLAP = 120

SUPPORTED = {".pdf", ".pptx", ".docx", ".txt", ".md"}


def _pdf(data: bytes):
    from pypdf import PdfReader

    for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, start=1):
        yield i, page.extract_text() or ""


def _pptx(data: bytes):
    from pptx import Presentation

    for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
        yield i, "\n".join(t for t in texts if t.strip())


def _docx(data: bytes):
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    # Word 문서는 페이지 정보가 없으므로 전체를 1페이지로 취급
    yield 1, "\n".join(parts)


def _text(data: bytes):
    yield 1, data.decode("utf-8", errors="replace")


_READERS = {".pdf": _pdf, ".pptx": _pptx, ".docx": _docx, ".txt": _text, ".md": _text}


def split_text(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = " ".join(text.split())
    if not text:
        return []
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        # 문장 경계에서 끊기 시도
        if end < len(text):
            cut = max(text.rfind(". ", start, end), text.rfind("다. ", start, end), text.rfind("\n", start, end))
            if cut > start + size // 2:
                end = cut + 1
        chunks.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in chunks if c]


def parse(filename: str, data: bytes) -> list[dict]:
    ext = PurePath(filename).suffix.lower()
    if ext not in _READERS:
        raise ValueError(f"지원하지 않는 형식: {ext}")
    return [
        {"page": page, "text": chunk}
        for page, text in _READERS[ext](data)
        for chunk in split_text(text)
    ]

import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "server"), str(ROOT / "client")]

import parsers  # noqa: E402
from store import Store  # noqa: E402


def test_split_text_overlaps_and_covers():
    text = "가나다라마. " * 400
    chunks = parsers.split_text(text, size=200, overlap=40)
    assert len(chunks) > 1
    assert all(len(c) <= 200 for c in chunks)


def test_parse_docx_and_pptx():
    from docx import Document
    from pptx import Presentation

    buf = io.BytesIO()
    d = Document(); d.add_paragraph("연차 휴가는 15일이다.")
    d.save(buf)
    assert "15일" in parsers.parse("a.docx", buf.getvalue())[0]["text"]

    buf = io.BytesIO()
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "2분기 매출"
    s = p.slides.add_slide(p.slide_layouts[1]); s.shapes.title.text = "3분기 목표"
    p.save(buf)
    chunks = parsers.parse("b.pptx", buf.getvalue())
    assert [c["page"] for c in chunks] == [1, 2]


def test_parse_rejects_unknown():
    with pytest.raises(ValueError):
        parsers.parse("x.exe", b"")


def test_store_search(tmp_path):
    st = Store(tmp_path / "i.db")
    st.upsert("a.pdf", 1.0, 10, [{"page": 1, "text": "A"}, {"page": 2, "text": "B"}], [[1, 0], [0, 1]])
    assert st.search([0.1, 0.9], k=1)[0]["text"] == "B"
    st.remove("a.pdf")
    assert st.search([1, 0]) == []


def test_server_requires_token(monkeypatch):
    from fastapi.testclient import TestClient
    import app as server_app

    monkeypatch.setattr(server_app, "API_TOKEN", "secret")
    c = TestClient(server_app.app)
    files = {"file": ("a.txt", "안녕하세요".encode())}
    assert c.post("/parse", files=files).status_code == 401
    r = c.post("/parse", files=files, headers={"Authorization": "Bearer secret"})
    assert r.json()["chunks"][0]["text"] == "안녕하세요"
